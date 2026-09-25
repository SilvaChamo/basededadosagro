import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Target Directory
target_dir = "/Users/imac/Desktop/APP/basededadosagro/Presentação"
img_dir = os.path.join(target_dir, "_images")

pie_img_path = os.path.join(img_dir, "grafico_operadores_pizza.png")
gantt_img_path = os.path.join(img_dir, "grafico_roteiro_gantt.png")
handshake_img_path = os.path.join(img_dir, "parceria_mader_maos_agro.png")

# Palette Colors
DARK_GREEN = RGBColor(5, 71, 54)     # #054736
EARTH_ORANGE = RGBColor(217, 107, 39) # #D96B27
DARK_SLATE = RGBColor(17, 24, 39)    # #111827
LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)    # White
BORDER_COLOR = RGBColor(226, 232, 240) # Slate 200
WHITE = RGBColor(255, 255, 255)

def create_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen aspect ratio (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # Helper: Set slide background color
    def set_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    # Helper: Add Standard Header
    def add_header(slide, title_text, category_text="BASE DE DADOS AGRO MOÇAMBIQUE"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = DARK_GREEN
        banner.line.fill.background()

        # Category / Antetitle
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = category_text.upper()
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = EARTH_ORANGE

        # Main Slide Title
        txBox_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.6))
        tf_title = txBox_title.text_frame
        p_title = tf_title.paragraphs[0]
        r_title = p_title.add_run()
        r_title.text = title_text
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(22)
        r_title.font.bold = True
        r_title.font.color.rgb = WHITE

    # Helper: Add Card Shape
    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================================================
    # SLIDE 1: CAPA INSTITUCIONAL
    # ==========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1, DARK_GREEN)

    # Decorative Orange Banner Accent
    acc = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    acc.fill.solid()
    acc.fill.fore_color.rgb = EARTH_ORANGE
    acc.line.fill.background()

    # Main Card
    add_card(slide1, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.2), bg_color=DARK_SLATE, border_color=EARTH_ORANGE)

    # Subtitle / Antetitle
    t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.6), Inches(10.5), Inches(0.5))
    p = t_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "PROPOSTA DE PARCERIA E MAPEAMENTO DIGITAL UNIFICADO"
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = EARTH_ORANGE

    # Title
    t_box2 = slide1.shapes.add_textbox(Inches(1.4), Inches(2.2), Inches(10.5), Inches(1.6))
    tf2 = t_box2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = "Base de Dados Agro — A Infraestrutura Digital do Agronegócio Moçambicano"
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(30)
    r2.font.bold = True
    r2.font.color.rgb = WHITE

    # Subtitle Details
    t_box3 = slide1.shapes.add_textbox(Inches(1.4), Inches(3.9), Inches(10.5), Inches(0.8))
    tf3 = t_box3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    r3 = p3.add_run()
    r3.text = "Apresentação Institucional e Oportunidades de Parceria Estratégica com o MADER / FAR, IP"
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(18)
    r3.font.color.rgb = RGBColor(226, 232, 240)

    # Presenter Meta Box
    t_box4 = slide1.shapes.add_textbox(Inches(1.4), Inches(5.1), Inches(10.5), Inches(1.0))
    tf4 = t_box4.text_frame
    p4 = tf4.paragraphs[0]
    r4 = p4.add_run()
    r4.text = "Apresentador: Silva Chamo (Director Geral — EntreCAMPOS)\nData: Setembro de 2026 | Maputo, Moçambique"
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(14)
    r4.font.bold = True
    r4.font.color.rgb = EARTH_ORANGE

    # ==========================================================================
    # SLIDE 2: O DESAFIO DO SECTOR AGRÁRIO MOÇAMBICANO
    # ==========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_bg(slide2, LIGHT_BG)
    add_header(slide2, "O Desafio do Sector: Fragmentação e Custos Operacionais")

    cards_data_s2 = [
        ("1. Mapeamento Fragmentado", 
         "Inexistência de um registo único, centralizado e pesquisável dos operadores do agronegócio, fornecedores de insumos e prestadores de serviços em Moçambique."),
        ("2. Barreira na Informação Fidedigna", 
         "Dificuldade na ligação rápida e transparente entre produtores rurais e fornecedores certificados de sementes e fertilizantes licenciados pelas Direcções Provinciais."),
        ("3. Elevados Encargos para o Estado", 
         "Necessidade constante do MADER e do FAR, IP de despender recursos a cada ciclo agrário para mapear beneficiários e fornecedores elegíveis no terreno.")
    ]

    for i, (title, desc) in enumerate(cards_data_s2):
        left = Inches(0.8 + i * 3.95)
        add_card(slide2, left, Inches(1.5), Inches(3.7), Inches(5.3))
        
        tb = slide2.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = DARK_GREEN

        tb_body = slide2.shapes.add_textbox(left + Inches(0.2), Inches(2.7), Inches(3.3), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        r_body = p_body.add_run()
        r_body.text = desc
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(13)
        r_body.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 3: DIMENSÃO ACTUAL E INDICADORES (COM IMAGEM PIZZA)
    # ==========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_bg(slide3, LIGHT_BG)
    add_header(slide3, "Dimensão Actual e Indicadores Chave do Sector")

    # Left Column: Macro Indicators Text Card
    add_card(slide3, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.5))
    tb_s3 = slide3.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_s3 = tb_s3.text_frame
    tf_s3.word_wrap = True

    p = tf_s3.paragraphs[0]
    r = p.add_run()
    r.text = "INDICADORES MACROECONÓMICOS OFICIAIS (INE/MADER)"
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = DARK_GREEN

    stats_list = [
        ("• Produção Agrária Nacional (2025/26): ", "20M+ de Toneladas (+5% de crescimento anual)."),
        ("• PIB Agrícola Nacional (Q2 2025): ", "91.4 Mil Milhões de MZN (~24% do PIB Nacional)."),
        ("• Emprego Agrícola: ", "70% da Força Laboral Nacional (Sustento da maioria da população)."),
        ("• Directório de Empresas: ", "950+ Empresas ativas mapeadas (~47% da amostra do sector privado comercial)."),
        ("• Estrutura Agrária: ", "4 Milhões de explorações familiares de subsistência vs. ~2.000 empresas comerciais.")
    ]

    for prefix, body in stats_list:
        p_item = tf_s3.add_paragraph()
        p_item.space_before = Pt(8)
        r_pre = p_item.add_run()
        r_pre.text = prefix
        r_pre.font.name = "Times New Roman"
        r_pre.font.bold = True
        r_pre.font.size = Pt(11.5)
        r_pre.font.color.rgb = DARK_SLATE
        
        r_b = p_item.add_run()
        r_b.text = body
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11.5)
        r_b.font.color.rgb = DARK_SLATE

    # Right Column: Embedded Pie Chart Graphic
    if os.path.exists(pie_img_path):
        slide3.shapes.add_picture(pie_img_path, Inches(6.7), Inches(1.4), width=Inches(5.8))

    # ==========================================================================
    # SLIDE 4: ALINHAMENTO ESTRATÉGICO COM PEDSA II, PNISA II E FACIM 2026
    # ==========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_bg(slide4, LIGHT_BG)
    add_header(slide4, "Alinhamento Estratégico com as Directrizes do Governo")

    pillars = [
        ("FACIM 2026 & Digitalização",
         "Alinhamento directo com o discurso do Presidente da República na abertura da 61.ª FACIM 2026, que exortou o uso da tecnologia como catalisador da produtividade no campo."),
        ("Apoio ao FAR, IP & Fomento",
         "Base de consulta fiável para constituição de listas de fornecedores e prestadores certificados para programas de fomento e subsídio (ex: Programa Sustenta)."),
        ("Extensionismo Digital",
         "Capacitação tecnológica dos extensionistas do MADER no terreno com tabelas digitais de insumos e diagnósticos fitossanitários computacionais.")
    ]

    for i, (title, desc) in enumerate(pillars):
        left = Inches(0.8 + i * 3.95)
        add_card(slide4, left, Inches(1.5), Inches(3.7), Inches(5.3))
        
        tb = slide4.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = EARTH_ORANGE

        tb_body = slide4.shapes.add_textbox(left + Inches(0.2), Inches(2.7), Inches(3.3), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        r_body = p_body.add_run()
        r_body.text = desc
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(13)
        r_body.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 5: INOVAÇÕES TECNOLÓGICAS, AGROBOTANICA AI E MÓDULO IIAM (ENRIQUECIDO)
    # ==========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_bg(slide5, LIGHT_BG)
    add_header(slide5, "Inovações Tecnológicas & Módulo de Telemetria IIAM", "INTELIGÊNCIA ARTIFICIAL E RASTREABILIDADE DE CAMPO")

    tools = [
        ("🌿 AgroBotanica AI & Scaneamento",
         "Identificação instantânea de pragas e análise do valor nutritivo das culturas pela câmara do telemóvel, funcionando online e offline no campo."),
        ("🔬 Módulo IIAM de Rastreabilidade",
         "Acompanhamento em tempo real do desempenho de sementes e variedades modificadas pelo IIAM no terreno, transmitindo telemetria foliar automática."),
        ("📍 Geolocalização de Insumos & Serviços",
         "Mapeamento georeferenciado de fornecedores licenciados por província e distrito, permitindo comparação de preços e insumos em tempo real.")
    ]

    for i, (title, desc) in enumerate(tools):
        left = Inches(0.8 + i * 3.95)
        add_card(slide5, left, Inches(1.5), Inches(3.7), Inches(5.3))
        
        tb = slide5.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Times New Roman"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = DARK_GREEN

        tb_body = slide5.shapes.add_textbox(left + Inches(0.2), Inches(2.7), Inches(3.3), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        r_body = p_body.add_run()
        r_body.text = desc
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(13)
        r_body.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 6: SEGURANÇA DE DADOS, ANTI-FRAUDE & VALIDAÇÃO (NOVO)
    # ==========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_bg(slide6, LIGHT_BG)
    add_header(slide6, "Segurança de Dados, Anti-Fraude & Validação em 3 Camadas", "GOVERNAÇÃO E INTEGRIDADE DO DIRECTÓRIO NACIONAL")

    cards_s6 = [
        ("1. Leitor Inteligente de Alvará (OCR)",
         "Upload obrigatório da licença/alvará comercial com validação automática por OCR de NUIT, Província, Razão Social e Data de Validade."),
        ("2. Consulta de Leitura com o Baú Único",
         "Integração de leitura (Read-Only API) com o Baú Único/Conservatória para confirmar o NUEL e eliminar a criação de empresas fantasmas."),
        ("3. Monetização & Triagem de Planos",
         "Plano Básico (1.000 MT/ano) com ativação prioritária e selo de Verificado. Cadastros gratuitos sujeitos a auditoria documental manual e física.")
    ]

    for i, (title, desc) in enumerate(cards_s6):
        left = Inches(0.8 + i * 3.95)
        add_card(slide6, left, Inches(1.5), Inches(3.7), Inches(5.3))
        
        tb = slide6.shapes.add_textbox(left + Inches(0.2), Inches(1.8), Inches(3.3), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Times New Roman"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = EARTH_ORANGE

        tb_body = slide6.shapes.add_textbox(left + Inches(0.2), Inches(2.7), Inches(3.3), Inches(3.8))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        p_body = tf_body.paragraphs[0]
        r_body = p_body.add_run()
        r_body.text = desc
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(13)
        r_body.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 7: O QUE PRETENDEMOS DO MADER — PROVEDORIA & CERTIFICAÇÃO
    # ==========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_bg(slide7, LIGHT_BG)
    add_header(slide7, "O Que Pretendemos do MADER: Provedoria & Certificação", "OBJECTIVOS ESTRATÉGICOS DA PARCERIA INSTITUCIONAL")

    # Top Highlight Box with Summarized Paragraph
    add_card(slide7, Inches(0.8), Inches(1.3), Inches(11.733), Inches(2.2), bg_color=DARK_GREEN, border_color=EARTH_ORANGE)
    tb_s7_top = slide7.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.333), Inches(2.0))
    tf_s7_top = tb_s7_top.text_frame
    tf_s7_top.word_wrap = True
    p_s7_top = tf_s7_top.paragraphs[0]
    r_s7_top = p_s7_top.add_run()
    r_s7_top.text = "“Mais do que uma simples ponte digital entre o Estado e o público alvo da cadeia de valor agrária, pretendemos que o MADER assuma a posição de Provedor Institucional da Plataforma Base de Dados Agro, chancelando-a e autenticando-a como o repositório digital oficial, certificado e de elevada confiança no país. Sob este enquadramento, cabe ao MADER impulsionar a abertura de novos canais de financiamento e programas de fomento em articulação com os seus parceiros de cooperação internacional, bem como promover e difundir activamente a utilização da plataforma junto dos seus aliados estratégicos e de todas as suas delegações provinciais e distritais a nível nacional.”"
    r_s7_top.font.name = "Times New Roman"
    r_s7_top.font.size = Pt(13)
    r_s7_top.font.bold = True
    r_s7_top.font.italic = True
    r_s7_top.font.color.rgb = WHITE

    # 4 Grid Cards below
    mader_pillars = [
        ("1. Conexão Estado-Público Alvo",
         "Ligar directamente o MADER aos agricultores, fornecedores de insumos, agrólogos, processadores e agentes do agronegócio."),
        ("2. Provedor & Chancela Oficial",
         "O MADER como entidade certificadora e autenticadora que chancela os dados como repositório digital fidedigno e aprovado."),
        ("3. Canais de Financiamento",
         "Abertura de linhas de crédito e programas de fomento em articulação com parceiros de cooperação internacional (FIDA, Banco Mundial, BAD)."),
        ("4. Difusão Nacional & Delegações",
         "Promoção activa junto de parceiros, aliados estratégicos e todas as delegações provinciais e distritais (DPAP e SDAE).")
    ]

    for i, (title, desc) in enumerate(mader_pillars):
        left = Inches(0.8 + i * 2.95)
        add_card(slide7, left, Inches(3.7), Inches(2.8), Inches(3.3))

        tb_p = slide7.shapes.add_textbox(left + Inches(0.15), Inches(3.85), Inches(2.5), Inches(0.8))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        p = tf_p.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Times New Roman"
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = EARTH_ORANGE

        tb_b = slide7.shapes.add_textbox(left + Inches(0.15), Inches(4.7), Inches(2.5), Inches(2.2))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        r_b = p_b.add_run()
        r_b.text = desc
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11.5)
        r_b.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 8: PROPOSTA DE PARCERIA E DECLARAÇÃO ESTRATÉGICA
    # ==========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_bg(slide8, LIGHT_BG)
    add_header(slide8, "Proposta de Parceria Estratégica MADER / FAR, IP")

    # Main Callout Box (Verde Escuro com Borda Laranja)
    add_card(slide8, Inches(1.0), Inches(1.5), Inches(11.333), Inches(3.0), bg_color=DARK_GREEN, border_color=EARTH_ORANGE)
    tb_s8 = slide8.shapes.add_textbox(Inches(1.3), Inches(1.7), Inches(10.733), Inches(2.6))
    tf_s8 = tb_s8.text_frame
    tf_s8.word_wrap = True
    p_s8 = tf_s8.paragraphs[0]
    r_s8 = p_s8.add_run()
    r_s8.text = "“Nós, na EntreCAMPOS, já fizemos o trabalho pesado de pesquisar e catalogar 950 das cerca de 2.000 empresas formais do país. Com a parceria institucional do MADER e do FAR, IP, pretendemos validar e auditar este universo inicial e alcançar as restantes 1.000 empresas, com o objectivo de dotar Moçambique do seu primeiro Mapa Digital com 100% dos operadores do agronegócio integrados.”"
    r_s8.font.name = "Times New Roman"
    r_s8.font.size = Pt(17)
    r_s8.font.bold = True
    r_s8.font.italic = True
    r_s8.font.color.rgb = WHITE

    # Bottom 2 Cards (Model & Difusão)
    add_card(slide8, Inches(1.0), Inches(4.8), Inches(5.5), Inches(2.2))
    tb_b1 = slide8.shapes.add_textbox(Inches(1.2), Inches(4.9), Inches(5.1), Inches(2.0))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p_b1 = tf_b1.paragraphs[0]
    r_b1 = p_b1.add_run()
    r_b1.text = "Modelo Win-Win Sem Custos Orçamentais\n"
    r_b1.font.bold = True
    r_b1.font.size = Pt(14)
    r_b1.font.color.rgb = DARK_GREEN
    r_b1_desc = p_b1.add_run()
    r_b1_desc.text = "A infraestrutura é cedida gratuitamente ao Estado em troca de validação institucional e atribuição do selo oficial de conformidade aos operadores."
    r_b1_desc.font.size = Pt(11.5)
    r_b1_desc.font.color.rgb = DARK_SLATE

    add_card(slide8, Inches(6.833), Inches(4.8), Inches(5.5), Inches(2.2))
    tb_b2 = slide8.shapes.add_textbox(Inches(7.033), Inches(4.9), Inches(5.1), Inches(2.0))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    p_b2 = tf_b2.paragraphs[0]
    r_b2 = p_b2.add_run()
    r_b2.text = "Canal de Comunicação e Difusão Massiva\n"
    r_b2.font.bold = True
    r_b2.font.size = Pt(14)
    r_b2.font.color.rgb = EARTH_ORANGE
    r_b2_desc = p_b2.add_run()
    r_b2_desc.text = "Difusão direta via Web, Email e SMS de editais, avisos fitossanitários, linhas de crédito e iniciativas governamentais a todos os agentes económicos."
    r_b2_desc.font.size = Pt(11.5)
    r_b2_desc.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 9: ROTEIRO DE COOPERAÇÃO E CRONOGRAMA (COM IMAGEM GANTT)
    # ==========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_bg(slide9, LIGHT_BG)
    add_header(slide9, "Roteiro de Cooperação e Cronograma de Implementação")

    # Embedded Gantt Chart Image
    if os.path.exists(gantt_img_path):
        slide9.shapes.add_picture(gantt_img_path, Inches(0.8), Inches(1.3), width=Inches(11.733))

    # Milestones Card below image
    add_card(slide9, Inches(0.8), Inches(4.8), Inches(11.733), Inches(2.2))
    tb_s9 = slide9.shapes.add_textbox(Inches(1.0), Inches(4.9), Inches(11.333), Inches(2.0))
    tf_s9 = tb_s9.text_frame
    tf_s9.word_wrap = True
    p_s9 = tf_s9.paragraphs[0]
    r_s9 = p_s9.add_run()
    r_s9.text = "FASES EXECUTIVAS DE IMPLEMENTAÇÃO DA PARCERIA:"
    r_s9.font.name = "Times New Roman"
    r_s9.font.bold = True
    r_s9.font.size = Pt(13)
    r_s9.font.color.rgb = DARK_GREEN

    milestones = [
        ("• Fase 1 (Outubro 2026): ", "Criação de Grupo de Trabalho Técnico com pontos focais da EntreCAMPOS, FAR, IP e MADER."),
        ("• Fase 2 (Outubro 2026): ", "Assinatura do Memorando de Entendimento (MOU de cooperação institucional)."),
        ("• Fase 3 (Nov–Dez 2026): ", "Projecto Piloto de Validação Provincial numa província prioritária (Manica ou Sofala)."),
        ("• Fase 4 (Janeiro 2027): ", "Lançamento Oficial e Escalamento para o Mapeamento 100% Nacional.")
    ]
    for prefix, body in milestones:
        p_m = tf_s9.add_paragraph()
        p_m.space_before = Pt(3)
        r_pre = p_m.add_run()
        r_pre.text = prefix
        r_pre.font.name = "Times New Roman"
        r_pre.font.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.color.rgb = EARTH_ORANGE
        r_b = p_m.add_run()
        r_b.text = body
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11)
        r_b.font.color.rgb = DARK_SLATE

    # ==========================================================================
    # SLIDE 10: AGRADECIMENTOS E CONTACTOS OFICIAIS
    # ==========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_bg(slide10, DARK_GREEN)

    add_card(slide10, Inches(1.5), Inches(1.2), Inches(10.333), Inches(5.2), bg_color=DARK_SLATE, border_color=EARTH_ORANGE)

    tb_s10 = slide10.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(4.5))
    tf_s10 = tb_s10.text_frame
    tf_s10.word_wrap = True

    p10 = tf_s10.paragraphs[0]
    r10 = p10.add_run()
    r10.text = "Juntos pela Transformação Digital do Agronegócio Moçambicano\n"
    r10.font.name = "Times New Roman"
    r10.font.size = Pt(24)
    r10.font.bold = True
    r10.font.color.rgb = WHITE

    p10_sub = tf_s10.add_paragraph()
    p10_sub.space_before = Pt(10)
    r10_sub = p10_sub.add_run()
    r10_sub.text = "A EntreCAMPOS reitera a sua inteira disponibilidade para colaborar de forma empenhada com o MADER e o FAR, IP na construção do ecossistema agrícola mais moderno, inclusivo e próspero da África Austral.\n"
    r10_sub.font.name = "Times New Roman"
    r10_sub.font.size = Pt(15)
    r10_sub.font.color.rgb = RGBColor(226, 232, 240)

    p10_meta = tf_s10.add_paragraph()
    p10_meta.space_before = Pt(20)
    r10_meta = p10_meta.add_run()
    r10_meta.text = (
        "Director Geral: Silva Chamo\n"
        "Telefone: +258 87 75 75 288\n"
        "E-mail: geral@basededadosagro.com / silva@visualdesignmoz.com\n"
        "Website Oficial: www.basededadosagro.com\n"
        "Maputo, Moçambique | Setembro de 2026"
    )
    r10_meta.font.name = "Times New Roman"
    r10_meta.font.size = Pt(14)
    r10_meta.font.bold = True
    r10_meta.font.color.rgb = EARTH_ORANGE

    # Save PowerPoint Presentation
    pptx_path = os.path.join(target_dir, "Base_de_Dados_Agro_Apresentacao_MADER_FARIP.pptx")
    prs.save(pptx_path)
    print(f"🎉 Apresentação em PowerPoint (.pptx) com 10 slides gerada com sucesso: {pptx_path}")

if __name__ == "__main__":
    create_presentation()
